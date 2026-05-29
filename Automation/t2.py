from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from openpyxl import load_workbook
import time


workbook = load_workbook("form_data1.xlsx")
sheet = workbook.active

driver = webdriver.Chrome()
driver.maximize_window()
wait = WebDriverWait(driver, 10)

for i in range(2, sheet.max_row + 1):
    first   = sheet.cell(row=i, column=1).value
    last    = sheet.cell(row=i, column=2).value
    email   = sheet.cell(row=i, column=3).value
    mobile  = sheet.cell(row=i, column=4).value
    gender  = sheet.cell(row=i, column=5).value
    subject = sheet.cell(row=i, column=7).value
    address = sheet.cell(row=i, column=8).value
    state   = sheet.cell(row=i, column=9).value
    city    = sheet.cell(row=i, column=10).value

    driver.get("https://demoqa.com/automation-practice-form")
    time.sleep(1)

    # Remove ads
    driver.execute_script("document.querySelectorAll('#fixedban,footer').forEach(e=>e.remove())")
    time.sleep(0.5)

    driver.find_element(By.ID, "firstName").send_keys(first)
    driver.find_element(By.ID, "lastName").send_keys(last)
    driver.find_element(By.ID, "userEmail").send_keys(email)
    driver.find_element(By.ID, "userNumber").send_keys(mobile)

    # Gender
    gender_map = {"Male":"gender-radio-1","Female":"gender-radio-2","Other":"gender-radio-3"}
    driver.execute_script("arguments[0].click()", driver.find_element(By.ID, gender_map[gender]))

    # Subject
    sub = driver.find_element(By.ID, "subjectsInput")
    sub.send_keys(subject)
    time.sleep(0.5)
    sub.send_keys(Keys.ENTER)

    driver.find_element(By.ID, "currentAddress").send_keys(address)

    # Scroll to state dropdown
    driver.execute_script("arguments[0].scrollIntoView()", driver.find_element(By.ID, "stateCity-wrapper"))
    time.sleep(0.5)

    # Remove ads again after scroll
    driver.execute_script("document.querySelectorAll('#fixedban,footer').forEach(e=>e.remove())")
    time.sleep(0.3)

    # State
    driver.find_element(By.XPATH, "//div[@id='state']//input").send_keys(state)
    time.sleep(0.5)
    state_option = driver.find_element(By.XPATH, f"//div[contains(@class,'option') and text()='{state}']")
    driver.execute_script("arguments[0].click()", state_option)
    time.sleep(0.4)

    # City
    driver.find_element(By.XPATH, "//div[@id='city']//input").send_keys(city)
    time.sleep(0.5)
    city_option = driver.find_element(By.XPATH, f"//div[contains(@class,'option') and text()='{city}']")
    driver.execute_script("arguments[0].click()", city_option)
    time.sleep(0.4)

    # Submit
    driver.execute_script("arguments[0].click()", driver.find_element(By.ID, "submit"))
    time.sleep(2)

    if "Thanks for submitting the form" in driver.page_source:
        print(f"PASS - Row {i}: {first} {last} submitted")
    else:
        print(f"FAIL - Row {i}: {first} {last} failed")

driver.quit()