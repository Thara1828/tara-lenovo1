from selenium import webdriver
from selenium.webdriver.common.by import By
from openpyxl import load_workbook
import time

# Load Excel
workbook = load_workbook("KDT.xlsx")
sheet = workbook.active

driver = None

# Loop through Excel rows
for i in range(2, sheet.max_row + 1):

    keyword = sheet.cell(row=i, column=2).value
    locator_type = sheet.cell(row=i, column=3).value
    locator_value = sheet.cell(row=i, column=4).value
    test_data = sheet.cell(row=i, column=5).value

    print("Executing Step:", keyword)

    if keyword == "open_browser":
        driver = webdriver.Chrome()
        driver.maximize_window()

    elif keyword == "open_url":
        driver.get(test_data)
        time.sleep(2)

    elif keyword == "enter_text":
        if locator_type == "id":
            driver.find_element(By.ID, locator_value).send_keys(test_data)
        elif locator_type == "xpath":
            driver.find_element(By.XPATH, locator_value).send_keys(test_data)

    elif keyword == "click":
        if locator_type == "id":
            driver.find_element(By.ID, locator_value).click()
        elif locator_type == "xpath":
            driver.find_element(By.XPATH, locator_value).click()

    elif keyword == "verify_text":
        element = None
        if locator_type == "xpath":
            element = driver.find_element(By.XPATH, locator_value)

        if test_data in element.text:
            print("Verification Passed")
        else:
            print("Verification Failed")

    elif keyword == "close_browser":
        driver.quit()

    time.sleep(2)