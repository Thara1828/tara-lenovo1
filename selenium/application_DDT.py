from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
import time
from openpyxl import load_workbook

workbook = load_workbook('login_data.xlsx')
sheet = workbook.active

driver = webdriver.Chrome()
driver.maximize_window()
driver.get("https://practicetestautomation.com/practice-test-login/")

rows = sheet.max_row

for i in range(2, rows + 1):
    username = sheet.cell(row=i, column=1).value
    password = sheet.cell(row=i, column=2).value

    driver.find_element(By.ID, "username").clear()
    driver.find_element(By.ID, "username").send_keys(username)

    driver.find_element(By.ID, "password").clear()
    driver.find_element(By.ID, "password").send_keys(password)

    driver.find_element(By.ID, "submit").click()
    time.sleep(3)

    if "Logged In Successfully" in driver.page_source:
        print("Testcase passed")
    else:
        print("Testcase failed")

    time.sleep(4)
    driver.get("https://practicetestautomation.com/practice-test-login/")

driver.quit()