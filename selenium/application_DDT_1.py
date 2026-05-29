from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
import time
from openpyxl import load_workbook

workbook = load_workbook('login_data(saucedemo).xlsx')
sheet = workbook.active

driver = webdriver.Chrome()
driver.maximize_window()
driver.get("https://www.saucedemo.com/")

rows = sheet.max_row

for i in range(2, rows + 1):
    username = sheet.cell(row=i, column=1).value
    password = sheet.cell(row=i, column=2).value

    driver.find_element(By.ID, "user-name").clear()
    driver.find_element(By.ID, "user-name").send_keys(username)

    driver.find_element(By.ID, "password").clear()
    driver.find_element(By.ID, "password").send_keys(password)

    driver.find_element(By.ID, "login-button").click()
    time.sleep(3)

    if "inventory" in driver.page_source:
        print("Testcase",i-1," passed")
    else:
        print("Testcase",i-1," failed")

    time.sleep(4)
    driver.get("https://www.saucedemo.com/")

driver.quit()